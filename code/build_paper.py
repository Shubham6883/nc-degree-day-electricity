"""Build the manuscript (.docx) from results.json and the figures."""
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
import json, math, os, re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = json.load(open(ROOT / 'results' / 'results.json'))
M, S, CP = R['models'], R['season'], R['comp']
FIG = str(ROOT / 'figures' / 'fig')
OLD = str(ROOT / 'figures')
B = R['bp']; EX = R['exclusion']; HT = R['heating']
OUT = str(ROOT / 'results' / 'Two Arms of One Curve.docx')


def p(x, d=3):
    if x < .001: return 'p < .001'
    return 'p = ' + f'{x:.{d}f}'.lstrip('0')
def nz(x, d=2):
    s = f'{x:.{d}f}'
    return s[1:] if s.startswith('0.') else ('-' + s[2:] if s.startswith('-0.') else s)
def num(x, d=2): return f'{x:.{d}f}'
def cm(x): return f'{round(x):,}'
def pv(x, d=3):
    return '< .001' if x < .001 else nz(x, d)
SUP = str.maketrans('0123456789', '⁰¹²³⁴⁵⁶⁷⁸⁹')
def sci(x, d=2):
    m_, e_ = f'{x:.{d}e}'.split('e')
    e_i = int(e_)
    return f'{m_} × 10' + ('⁻' if e_i < 0 else '') + str(abs(e_i)).translate(SUP)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m_ in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, m_, Inches(1))
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.line_spacing = 2.0
st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_ = hp.add_run()
for kind, val in [('fldChar', 'begin'), ('instrText', 'PAGE'), ('fldChar', 'separate'),
                  ('t', '1'), ('fldChar', 'end')]:
    e = OxmlElement(f'w:{kind}')
    if kind == 'fldChar': e.set(qn('w:fldCharType'), val)
    else: e.text = val; e.set(qn('xml:space'), 'preserve')
    r_._r.append(e)
hp.runs[0].font.name = 'Times New Roman'; hp.runs[0].font.size = Pt(12)


def para(text='', *, bold=False, italic=False, align='left', indent=True,
         space_after=0, spacing=2.0, keep=False):
    q = doc.add_paragraph()
    q.paragraph_format.line_spacing = spacing
    q.paragraph_format.space_after = Pt(space_after)
    q.paragraph_format.keep_with_next = keep
    q.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    if indent: q.paragraph_format.first_line_indent = Inches(0.5)
    if text:
        rr = q.add_run(text); rr.bold, rr.italic = bold, italic
    return q
def rich(parts, *, indent=True, spacing=2.0, space_after=0):
    q = doc.add_paragraph()
    q.paragraph_format.line_spacing = spacing
    q.paragraph_format.space_after = Pt(space_after)
    if indent: q.paragraph_format.first_line_indent = Inches(0.5)
    for t, s in parts:
        rr = q.add_run(t); rr.bold = 'b' in s; rr.italic = 'i' in s
    return q
def h1(t): para(t, bold=True, align='center', indent=False, keep=True)
def h2(t): para(t, bold=True, indent=False, keep=True)
def pagebreak(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def figure(n, title, path, note=None, width=6.0):
    para(f'Figure {n}', bold=True, indent=False, keep=True)
    para(title, italic=True, indent=False, keep=True)
    q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.line_spacing = 1.0
    q.paragraph_format.space_before = Pt(6); q.paragraph_format.space_after = Pt(6)
    q.add_run().add_picture(path, width=Inches(width))
    if note: rich([('Note. ', 'i'), (note, '')], indent=False, spacing=1.0, space_after=12)

def set_cell(cell, text, *, italic=False, align='left', size=10.5):
    cell.text = ''
    q = cell.paragraphs[0]
    q.paragraph_format.line_spacing = 1.0
    q.paragraph_format.space_after = Pt(3); q.paragraph_format.space_before = Pt(3)
    q.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    rr = q.add_run(text); rr.italic = italic
    rr.font.size = Pt(size); rr.font.name = 'Times New Roman'
def border(cell, edge, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    bs = tcPr.find(qn('w:tcBorders'))
    if bs is None: bs = OxmlElement('w:tcBorders'); tcPr.append(bs)
    el = bs.find(qn(f'w:{edge}'))
    if el is None: el = OxmlElement(f'w:{edge}'); bs.append(el)
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz)); el.set(qn('w:color'), '000000')

def table(n, title, header, rows, note=None, widths=None, aligns=None):
    para(f'Table {n}', bold=True, indent=False, keep=True)
    para(title, italic=True, indent=False, keep=True)
    nc = len(header)
    widths = widths or [6.5 / nc] * nc
    t = doc.add_table(rows=1 + len(rows), cols=nc)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    _lay = OxmlElement('w:tblLayout'); _lay.set(qn('w:type'), 'fixed')
    t._tbl.tblPr.append(_lay)
    _grid = t._tbl.find(qn('w:tblGrid'))
    if _grid is not None:
        for _gc, _w in zip(_grid.findall(qn('w:gridCol')), widths):
            _gc.set(qn('w:w'), str(int(round(_w * 1440))))
    aligns = aligns or (['left'] + ['center'] * (nc - 1))
    for j, hh in enumerate(header):
        c = t.rows[0].cells[j]
        set_cell(c, hh, align=aligns[j] if j else 'left')
        border(c, 'top'); border(c, 'bottom')
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            set_cell(c, v.lstrip('~'), italic=v.startswith('~'), align=aligns[j] if j else 'left')
            if i == len(rows): border(c, 'bottom')
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, c in enumerate(row.cells): c.width = Inches(widths[j])
    if note: rich([('Note. ', 'i'), (note, '')], indent=False, spacing=1.0, space_after=12)
    else: doc.add_paragraph().paragraph_format.space_after = Pt(12)

TITLE = ('Two Arms of One Curve: Cooling and Heating Degree Days and Monthly Residential '
         'Electricity Sales in North Carolina, 1990–2026')

for _ in range(4): para(indent=False)
para(TITLE, bold=True, align='center', indent=False)
para(indent=False)
para('Shubham Jalan', align='center', indent=False)
para('Independent Researcher', align='center', indent=False)
para('September 2026', align='center', indent=False)
pagebreak()

h1('Abstract')
para(
    'Degree-day studies of electricity demand in the Southeast usually focus on summer cooling, '
    'but North Carolina heats a large share of its homes with electricity. I estimated both the '
    'cooling and heating responses of monthly statewide residential electricity sales using '
    f'{int(R["N"])} months of public data (January 1990 to April 2026) from NOAA’s nClimDiv '
    'dataset and the U.S. Energy Information Administration’s Form EIA-861M. A model with cooling '
    f'degree days alone explained R² = {nz(M["M1"]["r2"])} of the variance in monthly sales. Adding '
    f'heating degree days raised R² to {nz(M["M2"]["r2"])}, and adding a linear time trend raised it '
    f'to {nz(M["M3"]["r2"])}. In the trend model, each cooling degree day was associated with '
    f'{cm(CP["M3"]["b_cdd"])} MWh of monthly sales, 95% CI '
    f'[{cm(M["M3"]["coef"]["CDD"]["lo"])}, {cm(M["M3"]["coef"]["CDD"]["hi"])}], and each heating '
    f'degree day with {cm(CP["M3"]["b_hdd"])} MWh, 95% CI '
    f'[{cm(M["M3"]["coef"]["HDD"]["lo"])}, {cm(M["M3"]["coef"]["HDD"]["hi"])}], with Newey–West '
    'standard errors. Because the state has more than twice as many heating as cooling degree days '
    'in a typical year, the heating share of annual sales came out larger than the cooling '
    f'share. The heating-to-cooling ratio was {num(CP["M3"]["ratio"])} with a linear trend and '
    f'{num(CP["M4"]["ratio"])} with calendar-month effects added, so the ordering is consistent but '
    'the size of the gap depends on the model. From January 2008 to April 2026, per-customer '
    f'sensitivity to cooling degree days fell by {abs(M["M7"]["pct_cdd"])*100:.0f}% to '
    f'{abs(M["M8"]["pct_cdd"])*100:.0f}% depending on the model, and heating sensitivity also fell. '
    'Whether cooling sensitivity fell faster than heating sensitivity depended on the model '
    f'({p(M["M7"]["contrast"]["p_norm"])} without calendar-month effects, '
    f'{p(M["M8"]["contrast"]["p_norm"])} with them). These are associations in monthly energy, not '
    'forecasts, and they say nothing about peak power.', indent=False)
para()
rich([('Keywords: ', 'i'),
      ('temperature response function, cooling degree days, heating degree days, degree-day '
       'regression, North Carolina, residential electricity demand', '')], indent=False)
pagebreak()

para(TITLE, bold=True, align='center', indent=False)
para(
    'Residential electricity demand in the southeastern United States is growing, and air '
    'conditioning is a big part of it. Nationally, air conditioning accounts for about 19% of '
    'residential electricity use (U.S. Energy Information Administration [EIA], 2024), and 93% of '
    'households in the South Census region have it (EIA, 2022). Duke Energy expects customer '
    'energy needs in the Carolinas to grow over the next 15 years at roughly eight times the rate '
    'of the previous 15 (Duke Energy, 2025).')
para(
    'The usual way to link temperature and demand is the degree day. A day’s cooling degree day '
    '(CDD) value is the number of degrees its mean temperature sits above 65 °F, heating degree '
    'days (HDD) count degrees below 65 °F, and monthly totals add up the daily values (EIA, '
    '2023b). Demand goes up when it gets either hotter or colder than the balance point, so plotting '
    'demand against temperature gives a curve with two arms: cooling on the right and heating on '
    'the left (Hu et al., 2024). Degree-day studies in the Southeast often look only at the cooling '
    'arm. That is reasonable in Florida or Texas, but it is a bigger assumption in a state where '
    'many homes heat with electricity.')
para(
    'North Carolina is one of those states. In the 2020 Residential Energy Consumption Survey, 38% of '
    'its housing units listed a central heat pump as their main heating equipment, behind South '
    'Carolina (41%) and Alabama (39%) but among the higher shares in the country (EIA, 2023a). '
    'That number counts one type of equipment and is not the share of homes '
    'heating with electricity, which the survey table does not report. At the same time, '
    'the state’s cooling degree days have been trending up and its heating degree days down '
    '(Kunkel et al., 2020), so the balance between the two arms may be shifting.')
para(
    'I asked three questions. First, how does the estimated temperature response change when '
    'heating degree days and a time trend are added to a cooling-only model? Second, how do the '
    'heating and cooling portions of annual sales compare, and how much does that comparison depend '
    'on the model? Third, has the response per customer stayed the same over the years for which '
    'customer counts are available? The data are observational statewide monthly totals, so the '
    'results are associations, not causal effects, and they are about monthly energy (MWh), not '
    'peak power.')

h1('Background')
h2('Degree-Day Models of Electricity Demand')
para(
    'Temperature is a strong predictor of electricity demand, though how strong depends on the '
    'region, season, and model. Cawthorne et al. (2021) modeled seasonal demand for balancing '
    'authorities in Tennessee and Texas. After removing the effect of population growth, they '
    'found that temperature explained 44% to 67% of seasonal demand variability (Figure 1). Their '
    'models used population-weighted temperature and population-adjusted demand, so that range is '
    'not a benchmark for the unadjusted state totals used here. Fache and Bhat (2024) applied '
    'degree days to Florida in a regression that also included population, employment, GDP, '
    'electricity price, and daylight hours, and found temperature to be among the strongest '
    'predictors of residential demand.')
figure(1, 'Residuals of Electricity Demand Versus Population-Weighted Average Temperature for '
          'Tennessee (A: Winter, B: Summer) and Texas (C: Winter, D: Summer)',
       f'{OLD}/fig01_cawthorne_2021.png',
       'From Cawthorne et al. (2021), Frontiers in Sustainable Cities, licensed under CC BY 4.0. '
       'Winter panels slope down and summer panels slope up, which is the two-armed response.')
para(
    'The idea behind this is the balance point: a building needs little heating or cooling near '
    'it, and demand rises as the temperature moves away from it in either direction (EIA, '
    '2023b). Hu et al. (2024) studied this temperature response function for 36 European countries '
    'and projected it to 2100 under different assumptions about insulation, heating '
    'electrification, air conditioning adoption, and passive cooling. In their scenarios, more air '
    'conditioning raises summer demand in hot regions, while better insulation and passive cooling '
    'offset part of the increase. Their work is a set of projections for Europe, not a measurement '
    'of past change in the United States, but it raises an obvious question for North Carolina: '
    'has the response curve here stayed put?')
figure(2, 'Illustrative Diagram of a Two-Armed Temperature Response Function',
       f'{FIG}02_conceptual.png',
       'Schematic drawn by the author from the balance-point concept (EIA, 2023b) and the '
       'temperature response function in Hu et al. (2024). The curves are illustrative and not '
       'fitted to data. The steeper and flatter cooling arms show what a change in the response '
       'would look like.', width=5.7)
h2('Methodological Considerations')
para(
    'Three issues shaped the design. First, monthly energy data are strongly seasonal and '
    'autocorrelated, and Granger and Newbold (1974) showed that regressing one trending series on '
    'another can make a relationship look far more significant than it is. Second, a cooling-only '
    'model leaves out the heating arm. Because heating and cooling degree days are tied together '
    'by the annual cycle, leaving one out changes the coefficient on the other. Third, household '
    'electricity use varies more in some months than others (Li et al., 2018), so comparisons '
    'should not assume equal variances. Li et al. studied 58 English households from July to '
    'December 2011, so their result shows the general pattern rather than anything specific to '
    'North Carolina.')
figure(3, 'Boxplots of Per-Household Mean Electrical Load (kW) for Monitored and Synthetic '
          'Households, July–December 2011', f'{OLD}/fig03_li_2018.png',
       'From Li et al. (2018), Energy and Buildings, licensed under CC BY 4.0.')
para(
    'Most published work covers Tennessee and Texas (Cawthorne et al., 2021), Florida (Fache & '
    'Bhat, 2024), or Europe (Colelli et al., 2023; Hu et al., 2024). I found few studies that '
    'estimate both arms for North Carolina at the monthly statewide level or check whether the arms '
    'have shifted over time. This study fills that gap using public federal data.')

h1('Method')
h2('Data Sources')
para(
    'Statewide monthly degree days and mean temperature came from NOAA’s nClimDiv dataset '
    '(National Oceanic and Atmospheric Administration, 2026): the cooling degree day '
    '(climdiv-cddcst), heating degree day (climdiv-hddcst), and mean temperature (climdiv-tmpcst) '
    'files, version 1.0.0 dated June 4, 2026, filtered to North Carolina (state code 031). Each '
    'record starts with a ten-character ID: characters 1–3 are the state, character 4 is the '
    'division (0 for statewide), characters 5–6 are the element, and characters 7–10 are the year. '
    'The three files use element codes 26, 25, and 02, matching their file names. Missing values '
    '(−9999) were recoded. nClimDiv statewide values are averaged by area, not by population.')
para(
    'Monthly residential electricity sales, customer counts, and average revenue per kilowatt-hour '
    'came from EIA Form EIA-861M (EIA, 2026), in two files covering January 1990 to December 2009 '
    'and January 2010 onward, filtered to North Carolina. Sales are in megawatt-hours. The two '
    'files have the same columns, no overlapping months, and no duplicate rows when combined.')
h2('Data Preparation and a Break in Customer Counts')
para(
    'I joined the climate and sales data on year and month. Every one of the '
    f'{int(R["N"])} North Carolina year–month records matched, with no duplicates, no gaps, and no '
    f'negative degree-day values, giving {int(R["N"])} consecutive months from January 1990 '
    f'through April 2026. Annual figures use only the {int(R["n_complete_years"])} complete '
    'calendar years (1990–2025), since 2026 has only four months.')
para(
    'Customer counts start in January 2007. I computed month-over-month percentage changes after '
    'checking that the customer series had no missing months. Six months changed by more than 2% '
    f'in absolute value (Table 1): a jump of {R["cust_jump_pct"]:+.1f}% in January 2008 and five '
    'months between June 2021 and February 2022 ranging from −2.79% to +3.22%. Leaving out January '
    f'2008, the largest change is {R["cust_other_max_pct"]:.2f}%. Leaving out the 2021–22 months '
    f'as well, it is {R["cust_excl_cluster_max"]:.2f}%.')
para(
    'These two anomalies are different. The 2021–22 changes flip between positive and negative and '
    f'add up to {R["cluster_net_pct"]:+.2f}% over May 2021 to April 2022, close to normal annual '
    'growth, so they bounce the series around without changing its level. January 2008 is a '
    'one-time jump that never reverses. It is far too large to be real customer growth, and a '
    'change in how EIA counted customers seems the most likely cause, but I did not find EIA '
    'documentation confirming this, so the cause is suspected rather than confirmed.')
para(
    'The decision to drop months before 2008 does not depend on the cause. In January 2008, sales '
    f'rose {R["jan08_yoy_sales"]:+.1f}% from a year earlier while the customer count rose '
    f'{R["jan08_yoy_cust"]:+.1f}%. Customer counts increased far more sharply than sales, so sales '
    'per customer drops suddenly at that point, and any per-customer trend fitted across the break '
    f'would be meaningless. All per-customer analyses therefore use the {int(R["n_pc"])} months '
    f'from January 2008 onward instead of all {int(R["n_pc_all"])} months from January 2007. I kept '
    'the 2021–22 months because they look unusual but not wrong, and the Results show what happens '
    'when they are removed. No data were dropped because of how they affected a test result.')
table(1, 'Month-Over-Month Changes in Reported Residential Customer Counts Exceeding 2%',
      ['Month', 'Customers', 'Change from prior month'],
      [[f'{a["year"]}-{a["month"]:02d}', cm(a['customers']), f'{a["mom"]:+.2f}%']
       for a in R['cust_anomalies']],
      note='Computed on the January 2007 to April 2026 customer series. These are the only months '
           'with an absolute change above 2%.',
      widths=[1.3, 1.6, 2.2], aligns=['left', 'right', 'right'])
figure(4, 'Residential Customer Counts Reported in EIA-861M, January 2007 to April 2026',
       f'{FIG}04_customer_break.png',
       'The series jumps between December 2007 and January 2008 and stays at the higher level. '
       'The 2021–22 changes go up and down and are kept. Shaded months before January 2008 are '
       'left out of all per-customer analyses.', width=6.1)
h2('Models')
para(
    'All tests used α = .05. I built the models up step by step so the effect of each addition '
    'to the cooling-only model could be seen directly. Table 2 lists every model. Models 1 through 5 '
    'predict monthly statewide sales, and Models 6 through 8 predict sales per customer over the '
    'shorter period with usable customer counts.')
table(2, 'Model Definitions',
      ['Model', 'Outcome', 'Predictors', 'Sample'],
      [['1', 'Monthly sales', 'CDD', f'{int(R["N"])} months'],
       ['2', 'Monthly sales', 'CDD, HDD', f'{int(R["N"])} months'],
       ['3', 'Monthly sales', 'CDD, HDD, centred linear year', f'{int(R["N"])} months'],
       ['4', 'Monthly sales', 'CDD, HDD, centred linear year, 11 calendar-month indicators',
        f'{int(R["N"])} months'],
       ['5', 'Monthly sales', 'CDD', f'{M["M5"]["n"]} May–September months'],
       ['6', 'Sales per customer', 'CDD, HDD, centred linear year', f'{int(R["n_pc"])} months'],
       ['7', 'Sales per customer', 'Model 6 plus CDD × year and HDD × year',
        f'{int(R["n_pc"])} months'],
       ['8', 'Sales per customer', 'Model 7 plus 11 calendar-month indicators',
        f'{int(R["n_pc"])} months']],
      note='CDD and HDD are monthly degree-day totals with a 65 °F base. The year term is centred '
           'on the mean of each sample: the 1990–2026 mean for Models 1–5 and '
           f'{R["pc_tbar"]:.3f} for Models 6–8. Centring does not change fitted values or '
           'interaction coefficients, but it does change what the main effects mean.',
      widths=[.7, 1.5, 2.85, 1.45], aligns=['center', 'left', 'left', 'left'])
para(
    'I report Models 3 and 4 side by side instead of choosing one, because they estimate the '
    'degree-day effects from different kinds of variation. Model 3 uses all month-to-month '
    'variation, including the annual cycle, so other seasonal factors that line up with weather '
    '(daylight, school calendars, holidays) can end up in the degree-day coefficients. Model 4 adds '
    'a separate intercept for each calendar month, so the degree-day effects come only from '
    'differences between years within the same month, such as a hot July compared with a mild '
    'one. Neither model is clearly right, and the true response is not guaranteed to fall between '
    'them. The gap between them shows how much the answer depends on how seasonality is handled.')
h2('Inference')
para(
    'Ordinary least squares standard errors assume independent errors, which monthly energy data '
    'do not have. Unless noted, every interval and p value for a time-series model uses Newey–West '
    'standard errors (Newey & West, 1987), which allow for autocorrelation and unequal variance, '
    'with a Bartlett kernel and a maximum lag of 12 months. Test statistics are compared with the '
    'standard normal distribution, and the t distribution result is also given where it would '
    'change a reported value. Table 4 shows ordinary and Newey–West (HAC) standard errors side by '
    'side.')
para(
    'I chose the 12-month lag because it covers one full seasonal cycle. It was not preregistered, '
    'so I also report results for lags from 3 to 24 months wherever a conclusion might depend on '
    'the choice. HAC standard errors fix the uncertainty estimates, but they do not fix a model '
    'that is missing something important.')
para(
    'Autocorrelation does not always make standard errors bigger. For Model 1 the HAC standard '
    'error is smaller than the ordinary one at a 12-month lag and larger at shorter lags, while for '
    'Models 3 and 4 it is larger at every lag. I also fit Models 1 and 3 with AR(1) feasible '
    'generalized least squares. A big difference between that estimate and the ordinary one means '
    'the result is sensitive to assumptions about the errors, though it does not say which '
    'estimate is closer to the truth.')
para(
    'Model 5 uses only May through September. Because those months are not consecutive, treating '
    'neighboring rows as one month apart would wrongly treat September and the next May as '
    'adjacent. For Model 5 I instead computed the Newey–West terms using the actual calendar '
    'distance between the retained months, so every pair of months one to twelve months apart '
    'counts, including pairs that straddle the autumn-to-spring gap (September to the next May is '
    'eight months). The Results compare this with two simpler alternatives. Model 5 is only included '
    'for comparison with earlier cooling-season studies.')
h2('Heating and Cooling Components')
para(
    'Multiplying a degree-day coefficient by the average number of degree days per year gives the '
    'part of average annual sales the model assigns to that term. These are model-based estimates, '
    'not metered end uses. I treated annual degree days as fixed, so the uncertainty shown covers '
    'only the coefficients, not the choice of model. The heating-to-cooling ratio depends on two '
    'coefficients, so I used the delta method with their full HAC covariance to get its standard '
    'error.')
para(
    'Since ratios can behave badly under the delta method, I also ran a moving-block bootstrap: '
    'blocks of 12 consecutive months were resampled with replacement until the original sample '
    'size was reached, 4,000 times with a fixed random seed, and the interval runs between the 2.5 '
    'and 97.5 percentiles. Because the blocks are put back together in random order, the time '
    'trend is no longer in calendar order within each resample. The bootstrap keeps short-range '
    'dependence but not the long-run trend, so it is a check on the delta-method interval rather '
    'than a separate estimate.')
h2('Diagnostics and Robustness Checks')
para(
    'Residual plots for Models 1, 3, and 4 are in Appendix Figures A1 to A3. Robustness checks '
    'include a quadratic time trend, dividing sales and degree days by the number of days in each '
    'month, dropping the last three months (EIA reports recent months as preliminary), correlations '
    'after detrending and differencing, and a comparison of total and per-customer sales over the '
    'same period.')

h1('Results')
h2('Descriptive Statistics')
para(
    f'The sample has {int(R["N"])} months from January 1990 through April 2026. Over the '
    f'{int(R["n_complete_years"])} complete years, North Carolina averaged {cm(R["ann_cdd"])} '
    f'cooling degree days and {cm(R["ann_hdd"])} heating degree days per year, a ratio of '
    f'{num(R["hdd_cdd_ratio"])} to 1. Monthly residential sales averaged {cm(R["mean_sales"])} MWh '
    f'(SD = {cm(R["sd_sales"])}). Table 3 reports the monthly variables, and Figure 5 plots sales '
    'against mean monthly temperature. Both arms show up clearly in the raw data.')
table(3, 'Descriptive Statistics for Monthly Climate and Electricity Variables, January 1990–April 2026',
      ['Variable', 'n', 'M', 'SD', 'Min', 'Max'],
      [['Residential sales (MWh)', f'{int(R["N"])}', cm(R['mean_sales']), cm(R['sd_sales']),
        cm(R['min_sales']), cm(R['max_sales'])],
       ['Cooling degree days', f'{int(R["N"])}', f'{R["mean_cdd"]:.1f}', f'{R["sd_cdd"]:.1f}', '0',
        f'{int(R["max_cdd"])}'],
       ['Heating degree days', f'{int(R["N"])}', f'{R["mean_hdd"]:.1f}', f'{R["sd_hdd"]:.1f}', '0',
        f'{int(R["max_hdd"])}'],
       ['Mean temperature (°F)', f'{int(R["N"])}', f'{R["mean_tavg"]:.1f}', f'{R["sd_tavg"]:.1f}',
        f'{R["min_tavg"]:.1f}', f'{R["max_tavg"]:.1f}'],
       ['~Mean monthly sales by season', '', '', '', '', ''],
       ['~   Winter (Dec–Feb)', f'{S["Winter"]["n"]}', cm(S['Winter']['mean']), cm(S['Winter']['sd']), '', ''],
       ['~   Spring (Mar–May)', f'{S["Spring"]["n"]}', cm(S['Spring']['mean']), cm(S['Spring']['sd']), '', ''],
       ['~   Summer (Jun–Aug)', f'{S["Summer"]["n"]}', cm(S['Summer']['mean']), cm(S['Summer']['sd']), '', ''],
       ['~   Fall (Sep–Nov)', f'{S["Fall"]["n"]}', cm(S['Fall']['mean']), cm(S['Fall']['sd']), '', '']],
      note='Degree days use a 65 °F base. Seasonal rows give the mean and standard deviation of '
           'monthly sales within each season.',
      widths=[2.45, .5, 1.1, 1.0, .9, .85],
      aligns=['left', 'center', 'right', 'right', 'right', 'right'])
figure(5, 'Monthly Residential Electricity Sales Against Mean Monthly Temperature',
       f'{FIG}05_response_function.png',
       'Each point is one month, colored by season. The solid curve is a LOWESS smoother. The '
       'dashed line marks the 65 °F degree-day base. Sales are in gigawatt-hours (1 GWh = 1,000 '
       'MWh).')
h2('Adding Heating Degree Days and a Trend')
para(
    f'Cooling degree days alone (Model 1) explained R² = {nz(M["M1"]["r2"])} of the variance in '
    f'monthly sales, with a slope of {cm(M["M1"]["coef"]["CDD"]["b"])} MWh per cooling degree day, '
    f'95% CI [{cm(M["M1"]["coef"]["CDD"]["lo"])}, {cm(M["M1"]["coef"]["CDD"]["hi"])}]. Figure 6 '
    f'shows why the fit is poor: the {int(R["zero_cdd_n"])} months with zero cooling degree days '
    'cover almost the whole range of sales, because a cooling variable cannot tell a mild November '
    'from a cold January.')
figure(6, 'Monthly Cooling Degree Days and Residential Electricity Sales',
       f'{FIG}06_cdd_scatter.png',
       'Points are colored by season. The line and band are the Model 1 OLS fit and its 95% OLS '
       'confidence band. Inference in the text uses HAC standard errors.')
para(
    f'Adding heating degree days raised R² from {nz(M["M1"]["r2"])} to {nz(M["M2"]["r2"])}, '
    f'adding a linear trend raised it to {nz(M["M3"]["r2"])}, and calendar-month effects raised it '
    f'to {nz(M["M4"]["r2"])} (Table 4). The cooling slope changed a lot along the way: the Model 3 '
    f'slope is {R["pct_increase_m1_to_m3"]:.0f}% higher than the Model 1 slope (equivalently, the '
    f'Model 1 slope is {R["pct_m1_below_m3"]:.0f}% lower). This fits what you would expect from '
    'leaving out heating. Months with no cooling degree days often have heavy heating loads, which '
    'pulls the Model 1 intercept up and flattens the cooling slope.')
para(
    'With AR(1) feasible GLS, the Model 1 cooling slope is '
    f'{cm(M["M1"]["gls"]["b_cdd"])} MWh per cooling degree day instead of '
    f'{cm(M["M1"]["coef"]["CDD"]["b"])}, with an estimated autocorrelation of '
    f'{num(M["M1"]["gls"]["rho"])}. For Model 3 the same method gives '
    f'{cm(M["M3"]["gls"]["b_cdd"])} instead of {cm(M["M3"]["coef"]["CDD"]["b"])}, a much smaller '
    'change. The cooling-only estimate depends heavily on assumptions about the errors, and the '
    'Model 3 estimate does not.')
LADDER = [('1', 'Sales ~ CDD', 'M1'),
          ('2', 'Sales ~ CDD + HDD', 'M2'),
          ('3', 'Sales ~ CDD + HDD + year', 'M3'),
          ('4', 'Sales ~ CDD + HDD + year + month effects', 'M4'),
          ('5', 'Cooling season: Sales ~ CDD', 'M5')]
def cell(k, v):
    c = M[k]['coef'].get(v)
    if c is None:
        return '—'
    if k == 'M5':
        return f"{cm(c['b'])}\n({cm(R['m5_ols_se'])})\n[{cm(R['m5_hac_cal_se'])}]"
    return f"{cm(c['b'])}\n({cm(c['ols_se'])})\n[{cm(c['se'])}]"
table(4, 'Regression Models Predicting Monthly Residential Electricity Sales',
      ['Model', 'Specification', 'CDD', 'HDD', 'Year', 'R²', 'DW'],
      [[num_, spec, cell(k, 'CDD'), cell(k, 'HDD'), cell(k, 'Year'),
        nz(M[k]['r2']), f'{M[k]["dw"]:.2f}']
       for num_, spec, k in LADDER],
      note=f'N = {int(R["N"])} months for Models 1–4 and {M["M5"]["n"]} May–September months '
           'for Model 5 (see Table 2). Entries are unstandardized coefficients in MWh per degree '
           'day, or MWh per year for the trend. OLS standard errors are in parentheses and HAC '
           'standard errors (Bartlett kernel, 12-month lag) in brackets. Model 5 HAC errors use '
           'actual calendar distance between months. All coefficients are significant at p < .001 '
           'under both. DW = Durbin–Watson statistic; values below 2 mean positive residual '
           'autocorrelation. Models 6–8, which use sales per customer over a shorter period, are '
           'in Table 8.',
      widths=[.65, 2.15, .85, .8, .9, .53, .52],
      aligns=['center', 'left', 'center', 'center', 'center', 'center', 'center'])
para(
    'Table 4 also shows that HAC and ordinary standard errors do not move together in a fixed way. '
    'For Model 1 the HAC standard error on the cooling slope is smaller than the ordinary one at a '
    f'12-month lag ({R["bw_sweep"]["M1"]["12"]:.0f} vs. {R["bw_sweep"]["M1"]["ols"]:.0f}) but '
    f'larger at shorter lags, reaching {R["bw_sweep"]["M1"]["3"]:.0f} at three months. For Model 3 '
    'the HAC standard error is larger than the ordinary one at every lag, from '
    f'{R["bw_sweep"]["M3"]["3"]:.0f} at three months to {R["bw_sweep"]["M3"]["24"]:.0f} at 24 '
    f'months, compared with an ordinary value of {R["bw_sweep"]["M3"]["ols"]:.0f}.')
figure(7, 'Added-Variable Plots for the Cooling and Heating Terms of Model 3',
       f'{FIG}07_two_arms.png',
       'Each panel plots sales against one degree-day variable after removing the effects of the '
       'other degree-day variable and the trend. Lines and bands are OLS fits with 95% OLS bands; '
       'the intervals printed in each panel are the HAC intervals reported in the text.',
       width=6.2)
h2('Heating and Cooling Components')
para(
    f'In Model 3, each cooling degree day is associated with {cm(CP["M3"]["b_cdd"])} MWh of '
    f'monthly sales, 95% CI [{cm(M["M3"]["coef"]["CDD"]["lo"])}, '
    f'{cm(M["M3"]["coef"]["CDD"]["hi"])}], and each heating degree day with '
    f'{cm(CP["M3"]["b_hdd"])} MWh, 95% CI [{cm(M["M3"]["coef"]["HDD"]["lo"])}, '
    f'{cm(M["M3"]["coef"]["HDD"]["hi"])}], so a cooling degree day counts about '
    f'{num(CP["M3"]["slope_ratio"])} times as much as a heating degree day. In Model 4 the slopes '
    f'are {cm(CP["M4"]["b_cdd"])} and {cm(CP["M4"]["b_hdd"])}, a ratio of '
    f'{num(CP["M4"]["slope_ratio"])}. Multiplying each slope by the average yearly degree days '
    'turns these into shares of annual sales (Table 5, Figure 8).')
table(5, 'Model-Implied Annual Components of Residential Sales Under Two Specifications',
      ['Quantity', 'Model 3', 'Model 4'],
      [['Cooling slope (MWh per CDD)', cm(CP['M3']['b_cdd']), cm(CP['M4']['b_cdd'])],
       ['Heating slope (MWh per HDD)', cm(CP['M3']['b_hdd']), cm(CP['M4']['b_hdd'])],
       ['Slope ratio (cooling ÷ heating)', num(CP['M3']['slope_ratio']), num(CP['M4']['slope_ratio'])],
       ['Cooling component (million MWh)', f'{CP["M3"]["e_cool"]/1e6:.1f}', f'{CP["M4"]["e_cool"]/1e6:.1f}'],
       ['Heating component (million MWh)', f'{CP["M3"]["e_heat"]/1e6:.1f}', f'{CP["M4"]["e_heat"]/1e6:.1f}'],
       ['Cooling share of annual sales', f'{CP["M3"]["share_cool"]:.1%}', f'{CP["M4"]["share_cool"]:.1%}'],
       ['Heating share of annual sales', f'{CP["M3"]["share_heat"]:.1%}', f'{CP["M4"]["share_heat"]:.1%}'],
       ['~Heating ÷ cooling component', f'~{num(CP["M3"]["ratio"])}', f'~{num(CP["M4"]["ratio"])}'],
       ['   HAC delta-method 95% CI',
        f'[{num(CP["M3"]["lo"])}, {num(CP["M3"]["hi"])}]',
        f'[{num(CP["M4"]["lo"])}, {num(CP["M4"]["hi"])}]'],
       ['   Moving-block bootstrap 95% CI',
        f'[{num(CP["M3"]["boot_lo"])}, {num(CP["M3"]["boot_hi"])}]',
        f'[{num(CP["M4"]["boot_lo"])}, {num(CP["M4"]["boot_hi"])}]']],
      note=f'Components are the coefficient times average yearly degree days over the '
           f'{int(R["n_complete_years"])} complete years ({cm(R["ann_cdd"])} CDD and '
           f'{cm(R["ann_hdd"])} HDD), compared with average annual sales of {cm(R["ann_sales"])} '
           'MWh. Intervals cover coefficient uncertainty only, not model choice. The delta-method '
           'interval uses the full HAC covariance of the two coefficients, which are strongly '
           f'correlated (HAC correlation {num(CP["M3"]["corr"])} in Model 3 and '
           f'{num(CP["M4"]["corr"])} in Model 4). The bootstrap uses 12-month blocks and 4,000 '
           'resamples.',
      widths=[2.9, 1.8, 1.8], aligns=['left', 'right', 'right'])
figure(8, 'Model-Implied Decomposition of Mean Annual Residential Sales Under Two Specifications',
       f'{FIG}08_decomposition.png',
       'Cooling and heating components are the degree-day coefficients times average yearly '
       'degree days. The remainder is the rest of average annual sales; it is what the model does '
       'not assign to weather, not a measurement of non-weather appliance use.', width=6.3)
para(
    'Both models put the heating component above the cooling component, but by very different '
    f'amounts: the ratio is {num(CP["M3"]["ratio"])} in Model 3, a gap of about '
    f'{abs(CP["M3"]["ratio"]-1)*100:.0f}%, and {num(CP["M4"]["ratio"])} in Model 4, a gap of about '
    f'{abs(CP["M4"]["ratio"]-1)*100:.0f}%. Neither interval in Table 5 captures this spread, because '
    'each one only reflects uncertainty within its own model. So I do not report a single ratio. '
    'The safe conclusion is that heating and cooling are roughly the same size, with heating at '
    'least as large. Three checks barely moved the Model 3 ratio: a quadratic trend gave '
    f'{num(R["sens"]["quadratic_trend"]["ratio"])}, dividing by days in the month gave '
    f'{num(R["sens"]["per_day"]["ratio"])}, and dropping the three preliminary months gave '
    f'{num(R["sens"]["drop_last3"]["ratio"])}. The difference between Models 3 and 4 comes from '
    'the calendar-month effects, not from the trend shape, month length, or preliminary data.')
h2('Seasonal Patterns')
para(
    'Table 6 compares average monthly sales between seasons. I regressed sales on season '
    'indicators and used the same 12-month Newey–West standard errors. Welch t tests allow unequal '
    'variances but still assume independent months, and the choice matters for one comparison: '
    'the Spring–Fall difference is not significant with Welch tests, '
    f'{p(R["seas"]["Spring-Fall"]["p_welch_bonf"])}, but is with HAC standard errors, '
    f'{p(R["seas"]["Spring-Fall"]["p_bonf"])}. Winter and Summer are not significantly different '
    'under either method (HAC interval '
    f'[{cm(R["seas"]["Winter-Summer"]["lo"])}, {cm(R["seas"]["Winter-Summer"]["hi"])}] MWh).')
table(6, 'Differences Between Seasonal Mean Monthly Sales, With Dependence-Robust Intervals',
      ['Comparison', 'M difference (MWh)', 'HAC 95% CI', 'p (HAC)', 'p (Welch)'],
      [[k.replace('-', ' − '), cm(v['diff']),
        f'[{cm(v["lo"])}, {cm(v["hi"])}]',
        ('1.00' if v['p_bonf'] > .995 else (nz(v['p_bonf'], 3) if v['p_bonf'] >= .001 else '< .001')),
        ('1.00' if v['p_welch_bonf'] > .995 else (nz(v['p_welch_bonf'], 3) if v['p_welch_bonf'] >= .001 else '< .001'))]
       for k, v in R['seas'].items()],
      note='HAC intervals and p values come from a regression of monthly sales on season '
           'indicators with Newey–West standard errors (Bartlett kernel, 12-month lag). Welch p '
           'values are from pairwise unequal-variance t tests. Both sets of p values are '
           'Bonferroni-corrected for six comparisons; the confidence intervals are not. Season '
           f'accounts for {nz(R["eta2"])} of the total variance in monthly sales (descriptive only).',
      widths=[1.5, 1.2, 2.15, .8, .8],
      aligns=['left', 'right', 'center', 'right', 'right'])
figure(9, 'Seasonal Distributions of Monthly Residential Sales and Pairwise Differences',
       f'{FIG}09_seasonal.png',
       'Left: boxes show the interquartile range, lines the median, and diamonds the mean. Right: '
       'differences in seasonal means with Newey–West 95% intervals, uncorrected for multiplicity; '
       'Bonferroni-corrected p values are in Table 6.', width=6.3)
para(
    f'Across {int(R["pair_n"])} years, average July sales ({cm(R["pair_jul"])} MWh) were '
    f'{cm(R["pair_diff"])} MWh higher than average January sales ({cm(R["pair_jan"])} MWh) '
    f'(SD of the difference = {cm(R["pair_sd"])}), which is not significant, '
    f't({int(R["pair_n"])-1}) = {num(R["pair_t"])}, {p(R["pair_p2"])}, 95% CI '
    f'[{cm(R["pair_ci95_lo"])}, {cm(R["pair_ci95_hi"])}]. The yearly differences show no real '
    f'autocorrelation (lag-1 r = {num(R["pair_lag1"])}, Durbin–Watson = {num(R["pair_dw"])}), so '
    'treating them as independent is reasonable.')
para(
    'A non-significant difference does not show the months are equal, so I also ran two one-sided '
    'equivalence tests (Lakens, 2017). With a margin of ±0.4 standard deviations of the paired '
    f'difference, equivalence held, {p(R["tost"]["0.4"]["pmax"])}; it also held at ±0.5 standard '
    f'deviations, {p(R["tost"]["0.5"]["pmax"])}, but not at ±0.3, '
    f'{p(R["tost"]["0.3"]["pmax"])}. This matches the 90% confidence interval, '
    f'[{cm(R["pair_ci90_lo"])}, {cm(R["pair_ci90_hi"])}] MWh, sitting inside the ±0.4 and ±0.5 '
    'bounds (Figure 10, Table 7). I picked these margins after seeing the data, so this result is '
    'exploratory: it shows roughly how big a difference the data can rule out. It applies to '
    'average monthly energy in two specific months, not to peak demand.')
table(7, 'Paired Comparison and Equivalence Tests, July Versus January Monthly Sales',
      ['Test', 'Bound (MWh)', 'p', 'Conclusion'],
      [[f'Paired t test, t({int(R["pair_n"])-1}) = {num(R["pair_t"])}', '—',
        nz(R['pair_p2'], 3), 'Fail to reject the null'],
       ['Equivalence test, ±0.3 SD', f'±{cm(R["tost"]["0.3"]["bound"])}',
        nz(R['tost']['0.3']['pmax'], 3), 'Equivalence not established'],
       ['Equivalence test, ±0.4 SD', f'±{cm(R["tost"]["0.4"]["bound"])}',
        nz(R['tost']['0.4']['pmax'], 3), 'Equivalence established'],
       ['Equivalence test, ±0.5 SD', f'±{cm(R["tost"]["0.5"]["bound"])}',
        nz(R['tost']['0.5']['pmax'], 3), 'Equivalence established']],
      note=f'n = {int(R["pair_n"])} years, 1990–2025. Mean difference = '
           f'{cm(R["pair_diff"])} MWh, standardized mean difference = {num(R["pair_dz"])}. '
           'Bounds are multiples of the standard deviation of the paired difference, chosen after '
           'seeing the data. For equivalence rows, p is the larger of the two one-sided p values.',
      widths=[2.6, 1.2, .75, 1.85], aligns=['left', 'right', 'right', 'left'])
figure(10, 'July Versus January Sales, Paired by Year, With Equivalence Bounds',
       f'{FIG}10_paired_equivalence.png',
       'Left: each line connects one year’s January and July sales; color shows which month was '
       'higher, and diamonds mark the means. Right: the mean paired difference with its 90% and '
       '95% confidence intervals, compared with the three equivalence margins.', width=6.3)
h2('Has the Per-Customer Response Changed?')
para(
    'Because customer counts start in 2007 and jump in January 2008, this section uses the '
    f'{int(R["n_pc"])} months from January 2008 to April 2026 (18 years and 4 months) with sales '
    'per customer as the outcome. The main estimates come from Models 7 and 8, which let both '
    'degree-day slopes change linearly over time. Model 6, without that interaction, is shown for '
    'comparison. Figure 11 also shows slopes fitted separately for each year, but each of those '
    'rests on only 12 months, so they are for illustration.')
para(
    'Sensitivity to cooling degree days declined in both models. The CDD × year interaction was '
    f'{sci(M["M7"]["coef"]["CDDxYear"]["b"])} per year in Model 7, 95% CI '
    f'[{sci(M["M7"]["coef"]["CDDxYear"]["lo"])}, {sci(M["M7"]["coef"]["CDDxYear"]["hi"])}], '
    f'{p(M["M7"]["coef"]["CDDxYear"]["p_norm"])}, and {sci(M["M8"]["coef"]["CDDxYear"]["b"])} per '
    f'year in Model 8, {p(M["M8"]["coef"]["CDDxYear"]["p_norm"])}. From the start to the end of the '
    f'period, that is a drop of {abs(M["M7"]["pct_cdd"])*100:.0f}% in Model 7 and '
    f'{abs(M["M8"]["pct_cdd"])*100:.0f}% in Model 8. One caution: in Model 7 the ordinary p value '
    f'is only {nz(M["M7"]["coef"]["CDDxYear"]["ols_p"], 3)}, while the HAC p value is '
    f'{nz(M["M7"]["coef"]["CDDxYear"]["p_norm"], 4)} (Table 8). In Model 8 both are well below '
    '.05. Across lags from 3 to 24 months, the HAC p value stays at or below '
    f'{max(v["p"] for k, v in R["bw_interaction"]["M7"].items() if k != "0"):.3f} in Model 7 and '
    'below .001 in Model 8.')
para(
    'The heating interaction was also negative: '
    f'{sci(M["M7"]["coef"]["HDDxYear"]["b"])} per year in Model 7, '
    f'{p(M["M7"]["coef"]["HDDxYear"]["p_norm"])}, and {sci(M["M8"]["coef"]["HDDxYear"]["b"])} in '
    f'Model 8, {p(M["M8"]["coef"]["HDDxYear"]["p_norm"])}. It is significant only in Model 8, but '
    'the estimate points the same direction in both, so the data do not suggest heating '
    'sensitivity stayed flat.')
para(
    'A separate question is whether cooling sensitivity fell faster, in percentage terms, than '
    'heating sensitivity. To compare them on the same scale, I divided each interaction by its own '
    'main effect and took the difference, g = β(CDD × year) ÷ β(CDD) − β(HDD × year) ÷ β(HDD), '
    'with a delta-method standard error from the full HAC covariance. Because the main effects are '
    'in the denominators, g depends on where the year variable is centred. It is centred at '
    f'{R["pc_tbar"]:.3f}, the middle of the period, so the comparison is made at the midpoint.')
para(
    f'In Model 7, g = {M["M7"]["contrast"]["b"]:+.5f} per year, HAC SE '
    f'{M["M7"]["contrast"]["se"]:.5f}, 95% CI [{M["M7"]["contrast"]["lo"]:+.5f}, '
    f'{M["M7"]["contrast"]["hi"]:+.5f}], {p(M["M7"]["contrast"]["p_norm"])} (normal reference; '
    f'{p(M["M7"]["contrast"]["p_t"])} with a t reference on {M["M7"]["contrast"]["dof"]} degrees of '
    f'freedom). In Model 8, g = {M["M8"]["contrast"]["b"]:+.5f} per year, HAC SE '
    f'{M["M8"]["contrast"]["se"]:.5f}, 95% CI [{M["M8"]["contrast"]["lo"]:+.5f}, '
    f'{M["M8"]["contrast"]["hi"]:+.5f}], {p(M["M8"]["contrast"]["p_norm"])} '
    f'({p(M["M8"]["contrast"]["p_t"])}). This is not driven by the lag choice: across lags from 3 '
    'to 24 months the Model 7 p value ranges from '
    f'{min(v["p"] for v in R["bw_contrast"]["M7"].values()):.3f} to '
    f'{max(v["p"] for v in R["bw_contrast"]["M7"].values()):.3f}, and the Model 8 p value from '
    f'{min(v["p"] for v in R["bw_contrast"]["M8"].values()):.3f} to '
    f'{max(v["p"] for v in R["bw_contrast"]["M8"].values()):.3f}. Whether cooling fell faster than '
    'heating depends on whether calendar-month effects are included, so I do not treat it as '
    'established.')
table(8, 'Per-Customer Degree-Day Models and the Proportional-Trend Contrast',
      ['Estimate', 'Model 6', 'Model 7', 'Model 8'],
      [['CDD main effect', sci(M['M6']['coef']['CDD']['b']), sci(M['M7']['coef']['CDD']['b']),
        sci(M['M8']['coef']['CDD']['b'])],
       ['HDD main effect', sci(M['M6']['coef']['HDD']['b']), sci(M['M7']['coef']['HDD']['b']),
        sci(M['M8']['coef']['HDD']['b'])],
       ['CDD × year', '—', sci(M['M7']['coef']['CDDxYear']['b']),
        sci(M['M8']['coef']['CDDxYear']['b'])],
       ['   OLS p', '—', pv(M['M7']['coef']['CDDxYear']['ols_p']),
        pv(M['M8']['coef']['CDDxYear']['ols_p'])],
       ['   HAC p (normal reference)', '—', pv(M['M7']['coef']['CDDxYear']['p_norm']), '< .001'],
       ['   Implied change over window', '—', f'{M["M7"]["pct_cdd"]*100:+.1f}%',
        f'{M["M8"]["pct_cdd"]*100:+.1f}%'],
       ['HDD × year', '—', sci(M['M7']['coef']['HDDxYear']['b']),
        sci(M['M8']['coef']['HDDxYear']['b'])],
       ['   OLS p', '—', pv(M['M7']['coef']['HDDxYear']['ols_p']),
        pv(M['M8']['coef']['HDDxYear']['ols_p'])],
       ['   HAC p (normal reference)', '—', pv(M['M7']['coef']['HDDxYear']['p_norm']),
        pv(M['M8']['coef']['HDDxYear']['p_norm'])],
       ['   Implied change over window', '—', f'{M["M7"]["pct_hdd"]*100:+.1f}%',
        f'{M["M8"]["pct_hdd"]*100:+.1f}%'],
       ['~Proportional-trend contrast g', '~—', f'~{M["M7"]["contrast"]["b"]:+.5f}',
        f'~{M["M8"]["contrast"]["b"]:+.5f}'],
       ['   HAC SE', '—', f'{M["M7"]["contrast"]["se"]:.5f}', f'{M["M8"]["contrast"]["se"]:.5f}'],
       ['   95% CI (normal reference)', '—',
        f'[{M["M7"]["contrast"]["lo"]:+.5f}, {M["M7"]["contrast"]["hi"]:+.5f}]',
        f'[{M["M8"]["contrast"]["lo"]:+.5f}, {M["M8"]["contrast"]["hi"]:+.5f}]'],
       ['   p (normal / t reference)', '—',
        f'{pv(M["M7"]["contrast"]["p_norm"])} / {pv(M["M7"]["contrast"]["p_t"])}',
        f'{pv(M["M8"]["contrast"]["p_norm"])} / {pv(M["M8"]["contrast"]["p_t"])}'],
       ['Model R²', nz(M['M6']['r2']), nz(M['M7']['r2']), nz(M['M8']['r2'])],
       ['Durbin–Watson', f'{M["M6"]["dw"]:.2f}', f'{M["M7"]["dw"]:.2f}', f'{M["M8"]["dw"]:.2f}']],
      note=f'n = {int(R["n_pc"])} months, January 2008 to April 2026. The outcome is sales per '
           'residential customer in MWh. Main effects are MWh per customer per degree day, and '
           'interactions are MWh per customer per degree day per year. The year variable is '
           f'centred at {R["pc_tbar"]:.3f}. Implied change compares the fitted slope at the first '
           'and last months. HAC standard errors use a Bartlett kernel with a 12-month lag.',
      widths=[2.15, 1.3, 1.5, 1.55], aligns=['left', 'right', 'right', 'right'])
figure(11, 'Annual Per-Customer Response Slopes and the Proportional-Trend Contrast',
       f'{FIG}11_stability.png',
       'Left and centre: each point is a slope from a regression of sales per customer on cooling '
       'and heating degree days fitted within one year, with 95% OLS error bars from that year’s '
       '12 months; lines and bands are OLS trends through those slopes. Right: the contrast g from '
       'each pooled model with its Newey–West 95% interval. Estimates in the text come from the '
       'pooled models in Table 8.', width=6.4)
para(
    'Table 9 shows how these results change with the time window. Dropping the five 2021–22 '
    'months from Table 1 changes almost nothing (Model 8 contrast '
    f'{p(EX["2008-2026 less 2021-22 anomalies"]["fe"]["p_contrast"])}, Model 7 '
    f'{p(EX["2008-2026 less 2021-22 anomalies"]["nofe"]["p_contrast"])}). Starting in January 2009 '
    'keeps the cooling decline significant in both models but weakens the Model 8 contrast to '
    f'{p(EX["2009-2026"]["fe"]["p_contrast"])}. Including 2007, which I excluded because of the '
    'customer-count break, makes the cooling decline look larger but makes both contrasts clearly '
    f'non-significant ({p(EX["2007-2026 (2007 retained)"]["fe"]["p_contrast"])} with '
    'calendar-month effects). The cooling decline shows up in every window. The contrast is '
    'significant only with calendar-month effects and only for windows starting in 2008.')
table(9, 'Sensitivity of the Per-Customer Results to the Estimation Window',
      ['Window', 'n', 'CDD × year p', 'Implied cooling change', 'Contrast p'],
      [[lab, str(v['fe']['n']),
        f'{pv(v["nofe"]["p_cddxyear"])} / {pv(v["fe"]["p_cddxyear"])}',
        f'{v["nofe"]["pct_cdd"]*100:+.1f}% / {v["fe"]["pct_cdd"]*100:+.1f}%',
        f'{pv(v["nofe"]["p_contrast"])} / {pv(v["fe"]["p_contrast"])}']
       for lab, v in EX.items()],
      note='Each cell gives the value without calendar-month effects (Model 7) and then with them '
           '(Model 8). All p values are HAC with a 12-month lag and a normal reference. The first '
           'row is the main analysis. Windows were not chosen based on their results.',
      widths=[2.0, .45, 1.25, 1.55, 1.2],
      aligns=['left', 'center', 'right', 'right', 'right'])
h2('Robustness')
para(
    'Removing a linear trend from both series raised the CDD–sales correlation from '
    f'{nz(R["r_all"])} to {nz(R["r_detrend"])}. The correlation was {nz(R["r_diff1"])} on first '
    f'differences and {nz(R["r_diff12"])} on twelve-month differences. The relationship holds up '
    'after each transformation, which argues against it being a spurious trend correlation of '
    'the kind Granger and Newbold (1974) described, though it does not rule out every shared '
    'influence.')
para(
    'Over the same January 2008 to April 2026 window, the CDD correlation was '
    f'{nz(R["match_total_r"])} for total sales and {nz(R["match_pc_r"])} for sales per customer, '
    'so dividing by customers does not change the strength of the relationship on its own. Heating '
    f'degree days correlated with winter sales at {nz(R["win_hdd_r"])} and cooling degree days '
    f'with summer sales at {nz(R["sum_cdd_r"])}. Those are computed on different months, so they '
    'cannot be used to say which arm is stronger; the joint models handle that comparison.')
para(
    'Restricting Model 1 to May through September raised R² to '
    f'{nz(M["M5"]["r2"])}, with a slope of {cm(R["m5_slope"])} MWh per cooling degree day. Using '
    'only summer months does not make the data independent. The residual autocorrelation between '
    f'truly adjacent months within each May–September block ({int(R["cool_lag1_npairs"])} pairs) '
    f'was {nz(R["cool_lag1_within"])}. Computing it naively across consecutive rows gives '
    f'{nz(R["cool_lag1_naive"])}, because that treats September and the next May as neighbors. The '
    'same issue affects the Newey–West standard error for the Model 5 slope. The naive estimate is '
    f'{num(R["m5_hac_naive_se"], 1)}. Using only pairs within a single summer gives '
    f'{num(R["m5_hac_block_se"], 1)}, which avoids the mistake but throws away real pairs less '
    'than twelve months apart across the year boundary. Using actual calendar distance, so all '
    f'{int(R["m5_pairs_total"])} eligible pairs count ({int(R["m5_pairs_cross"])} across years and '
    f'{int(R["m5_pairs_within"])} within years), gives {num(R["m5_hac_cal_se"], 1)}, compared with '
    f'an ordinary standard error of {num(R["m5_ols_se"], 1)}. The calendar-distance figure is the '
    'one reported in Table 4.')
para(
    'Residual diagnostics are in Appendix Figures A1 to A3. Model 1 residuals are strongly '
    f'autocorrelated (lag-1 {nz(R["acf1_M1"])}, lag-12 {nz(R["acf12_M1"])}, Durbin–Watson '
    f'{num(M["M1"]["dw"])}). Model 3 lowers the lag-1 value to {nz(R["acf1_M3"])} but still has '
    f'{nz(R["acf12_M3"])} at lag 12, and Model 4 lowers the lag-12 value to {nz(R["acf12_M4"])}. '
    'Some seasonal dependence remains in every model, which is why I use HAC standard errors '
    'throughout.')
para(
    'The Breusch–Pagan test depends on which variables go into the auxiliary regression, so I '
    f'report more than one variant. For Model 1 with CDD, LM = {num(B["M1, aux on CDD"]["lm"])} on 1 '
    f'degree of freedom, {p(B["M1, aux on CDD"]["p"], 2)}. For Model 3 with CDD and HDD, '
    f'LM = {num(B["M3, aux on CDD+HDD"]["lm"])} on 2 degrees of freedom, '
    f'{p(B["M3, aux on CDD+HDD"]["p"])}; adding the trend, which matches the fitted model, gives '
    f'LM = {num(B["M3, aux on CDD+HDD+trend"]["lm"])} on 3 degrees of freedom, '
    f'{p(B["M3, aux on CDD+HDD+trend"]["p"])}. These tests assume independent errors, so they are '
    'only a rough guide. The HAC standard errors already allow for unequal variance.')
figure(12, 'Monthly Residential Electricity Sales, January 1990 to April 2026',
       f'{FIG}12_timeseries.png',
       'The curve is a LOWESS trend. Sales peak twice each year, in winter and summer, on top of '
       'long-run growth.',
       width=6.4)

h1('Discussion')
para(
    'A cooling-only degree-day model fits North Carolina’s residential electricity sales poorly, '
    'and the reason is that it leaves out heating, not that the cooling relationship is weak. '
    f'Adding heating degree days raised R² from {nz(M["M1"]["r2"])} to {nz(M["M2"]["r2"])}, and '
    f'adding a linear trend raised it to {nz(M["M3"]["r2"])}. The second jump comes from long-run '
    'growth, not weather. Leaving out heating also distorts the cooling coefficient itself: the '
    f'Model 3 estimate is {R["pct_increase_m1_to_m3"]:.0f}% higher than the Model 1 estimate. In a '
    'state where many homes heat with electricity, the heating arm cannot just be set aside as '
    'outside the scope of a study.')
para(
    'The comparison of heating and cooling components is the result that depends most on the '
    'model. With a linear trend, heating is larger than cooling by about '
    f'{abs(CP["M3"]["ratio"]-1)*100:.0f}%; with calendar-month effects, the gap grows to about '
    f'{abs(CP["M4"]["ratio"]-1)*100:.0f}%. Both models are reasonable, and they use different '
    'variation, so the honest summary is an ordering plus a range. Weather-driven heating and '
    'cooling are about the same size in North Carolina, each accounting for roughly 15% to 23% of '
    'annual residential sales depending on the model, with heating the larger.')
para(
    'The slope ratio should not be read as a statement about equipment efficiency. Air '
    'conditioners and heat pumps both use the same vapor-compression cycle and both move more heat '
    'than the electricity they use, so efficiency alone does not explain why a cooling degree day '
    'is tied to more electricity than a heating degree day. A more likely explanation is which '
    'homes respond: nearly all homes in the South cool with electricity (EIA, 2022), but only some '
    'heat with it, and homes that heat with gas or oil still use some electricity for blowers, '
    'pumps, and controls. I could not put numbers on this. The 2020 Residential Energy Consumption '
    'Survey reports main heating equipment, not main heating fuel: '
    f'{HT["nc_central_heat_pump_pct"]}% of North Carolina housing units use a central heat pump and '
    f'{HT["nc_furnace_pct"]}% use a furnace, but the table does not split furnaces by fuel or show '
    'electric resistance heating separately (EIA, 2023a). So the share of homes heating with '
    'electricity is larger than the heat-pump share by an unknown amount. Other possible factors '
    'include a balance point that is not exactly 65 °F, humidity loads in summer that degree days '
    'do not capture, and heat pumps losing efficiency and switching to resistance backup in cold '
    'weather. This study measures the overall association and cannot separate these.')
para(
    'For context, the same survey puts North Carolina’s central heat-pump share '
    f'({HT["nc_central_heat_pump_pct"]}%) below South Carolina ({HT["sc"]}%) and Alabama '
    f'({HT["al"]}%) and above Tennessee ({HT["tn"]}%) and Florida ({HT["fl"]}%). Differences this '
    'small may not be statistically significant in the survey.')
para(
    'The seasonal results support describing the year as having two peaks, but not much more. '
    'July and January average sales are close enough to be equivalent within an exploratory '
    'margin, and winter and summer averages are not significantly different. These are statements '
    'about monthly energy. Two months with similar megawatt-hours can have very different peak '
    'megawatt demands, so nothing here says which season is harder on the grid. Statewide '
    'residential sales also do not represent the full demand of any one utility.')
para(
    'The stability results contain three separate claims. First, per-customer cooling sensitivity '
    'declined from 2008 to 2026. This held in both models, at every lag, and in every time window '
    'I tried. Second, per-customer heating sensitivity also declined. The estimate is negative in '
    'both models and significant once calendar-month effects are included. Third, whether cooling '
    'declined faster than heating is not settled, because the answer depends on whether '
    'calendar-month effects are included, and the data cannot decide between those models. So the '
    'supported finding is a general decline in per-customer weather sensitivity.')
para(
    'I did not try to explain the decline. Possible causes include newer, more efficient heating '
    'and cooling equipment (federal minimum standards for central air conditioners rose in 2006, '
    'again for southern states in 2015, and again with the SEER2 test procedure in 2023; '
    'Air-Conditioning, Heating, and Refrigeration Institute, 2023), better-insulated buildings, '
    'changes in household or home size, response to rising prices, and changes in who the '
    'customers are as the state grows. This study measures none of these. Average revenue per '
    'kilowatt-hour, the only price-like variable available, is revenue divided by sales, so it is '
    'tied to the outcome itself and would not cleanly control for price.')
para(
    'A falling per-customer coefficient does not mean electricity use is falling. The number of '
    f'customers grew from {cm(R["cust_first"])} in January 2008 to {cm(R["cust_last"])} in April '
    '2026, so total weather-driven use can rise even while each customer responds less to each '
    'degree day.')

h1('Conclusion')
para(
    'Cooling degree days are positively associated with monthly residential electricity sales in '
    'North Carolina, but a model built on them alone misses half the picture. Adding heating '
    f'degree days raises R² from {nz(M["M1"]["r2"])} to {nz(M["M2"]["r2"])}, a linear trend raises '
    f'it to {nz(M["M3"]["r2"])}, and the cooling coefficient changes a lot along the way. Heating '
    'and cooling account for similar shares of annual sales, with heating larger in every model, '
    f'though the gap ranges from about {abs(CP["M3"]["ratio"]-1)*100:.0f}% to about '
    f'{abs(CP["M4"]["ratio"]-1)*100:.0f}% depending on how calendar months are handled. '
    'Per-customer sensitivity to degree days declined between 2008 and 2026, clearly for cooling '
    'and less clearly for heating. Whether cooling declined faster depends on the model. These '
    'are in-sample associations from statewide monthly totals, not forecasts, and they do not '
    'address peak power.')
h2('Limitations')
para(
    'This is an observational study of statewide totals, so it cannot see county-level '
    'differences in climate, housing, or rates, and it does not identify causal effects. nClimDiv '
    'degree days are weighted by area, not population, so they do not perfectly match where '
    'electricity is used. This is a kind of measurement error, but because it is systematic I '
    'cannot say which way, if any, it pushes the slopes.')
para(
    'Some seasonal autocorrelation remains in every model. HAC standard errors account for it in '
    'the inference, but I did not fit a model of the residual structure itself. The linear trend '
    'in Model 3 stands in for customer growth, income, prices, efficiency, and building changes '
    'all at once. The per-customer analysis starts in 2008 because of the customer-count break, '
    'whose cause is suspected rather than documented. The 2021–22 customer-count anomalies are '
    'unexplained, though removing them does not change the results.')
para(
    'EIA revises recent months, and I could not compare the last months of this extract with a '
    'later release, although dropping the last three months leaves the Model 3 ratio at '
    f'{num(R["sens"]["drop_last3"]["ratio"])}. I identified state code 031 as North Carolina from '
    'the file layout and did not check it against a separate NCEI code table. Winter and Spring '
    'each include one more year than Summer and Fall; using only complete years shifts the Winter '
    f'mean by {R["season_mean_complete_Winter"] - R["season"]["Winter"]["mean"]:+,.0f} MWh and '
    f'leaves the share of variance explained by season at {nz(R["eta2_complete"])} '
    f'(vs. {nz(R["eta2"])}).')
para(
    'All R² values are in-sample fit to the months used for estimation, and the paper makes no '
    'forecasting claim. Testing forecasts would need held-out later months, rolling-origin '
    'evaluation, comparison with seasonal-naïve and trend-only baselines, and a distinction between '
    'observed and forecast degree days.')
h2('Future Research')
para(
    'Four extensions follow. A regression with seasonal autoregressive errors would model the '
    'leftover dependence directly instead of only correcting the standard errors. County-level or '
    'balancing-authority data would allow population-weighted degree days. Hourly or daily load '
    'data would make it possible to study peak demand. Repeating the stability analysis across '
    'southern states with different electric-heating shares, using heating-fuel data, would help '
    'separate the possible explanations listed above.')

h2('Data and Code Availability')
para(
    'All data are public. Degree days and mean temperature come from the NOAA nClimDiv statewide '
    'files climdiv-cddcst, climdiv-hddcst, and climdiv-tmpcst (version 1.0.0, June 4, 2026), '
    'filtered to state code 031. Electricity sales, customer counts, and average revenue per '
    'kilowatt-hour come from EIA Form EIA-861M files for January 1990 to December 2009 and January '
    '2010 onward, filtered to North Carolina residential sales. Both were downloaded in June 2026.')
para(
    f'The merged data have {int(R["N"])} consecutive months. Customer counts are available for '
    f'{int(R["n_pc_all"])} months from January 2007, of which the {int(R["n_pc"])} months from '
    f'January 2008 are used. Annual figures use the {int(R["n_complete_years"])} complete years.')
para(
    f'The analysis was run in Python {R["software"]["python"]} with NumPy '
    f'{R["software"]["numpy"]} and pandas {R["software"]["pandas"]}. Newey–West standard errors '
    'and the t, F, and chi-square p values are computed in the project’s own code, which is '
    'checked against published critical values. The code, data, analysis panel, results file, '
    'and logs are available at https://github.com/Shubham6883/nc-degree-day-electricity.')
pagebreak()
h1('References')
refs = [
 'Air-Conditioning, Heating, and Refrigeration Institute. (2023). <i>2023 energy efficiency '
 'standards.</i> https://www.ahrinet.org/2023-energy-efficiency-standards',
 'Cawthorne, D., de Queiroz, A. R., Eshraghi, H., Sankarasubramanian, A., & DeCarolis, J. F. '
 '(2021). The role of temperature variability on seasonal electricity demand in the Southern US. '
 '<i>Frontiers in Sustainable Cities, 3</i>, Article 644789. '
 'https://doi.org/10.3389/frsc.2021.644789',
 'Colelli, F. P., Wing, I. S., & De Cian, E. (2023). Air-conditioning adoption and electricity '
 'demand highlight climate change mitigation–adaptation tradeoffs. <i>Scientific Reports, 13</i>(1), '
 'Article 4413. https://doi.org/10.1038/s41598-023-31469-z',
 'Duke Energy. (2025, October 1). <i>Duke Energy files 2025 Carolinas Resource Plan, continues '
 'modernizing energy infrastructure to support future growth</i> [Press release]. '
 'https://news.duke-energy.com/releases/duke-energy-files-2025-carolinas-resource-plan-continues-'
 'modernizing-energy-infrastructure-to-support-future-growth',
 'Fache, A., & Bhat, M. G. (2024). Temperature sensitive electricity demand and policy implications '
 'for energy transition: A case study of Florida, USA. <i>Frontiers in Sustainable Energy Policy, '
 '2</i>, Article 1271035. https://doi.org/10.3389/fsuep.2023.1271035',
 'Granger, C. W. J., & Newbold, P. (1974). Spurious regressions in econometrics. <i>Journal of '
 'Econometrics, 2</i>(2), 111–120. https://doi.org/10.1016/0304-4076(74)90034-7',
 'Hu, W., Scholz, Y., Yeligeti, M., Deng, Y., & Jochem, P. (2024). Future electricity demand for '
 'Europe: Unraveling the dynamics of the Temperature Response Function. <i>Applied Energy, 368</i>, '
 'Article 123387. https://doi.org/10.1016/j.apenergy.2024.123387',
 'Kunkel, K. E., Easterling, D. R., Ballinger, A., Bililign, S., Champion, S. M., Corbett, D. R., '
 'Dello, K. D., Dissen, J., Lackmann, G. M., Luettich, R. A., Jr., Perry, L. B., Robinson, W. A., '
 'Stevens, L. E., Stewart, B. C., & Terando, A. J. (2020). <i>North Carolina climate science '
 'report.</i> North Carolina Institute for Climate Studies. https://ncics.org/nccsr',
 'Lakens, D. (2017). Equivalence tests: A practical primer for t tests, correlations, and '
 'meta-analyses. <i>Social Psychological and Personality Science, 8</i>(4), 355–362. '
 'https://doi.org/10.1177/1948550617697177',
 'Li, M., Allinson, D., & He, M. (2018). Seasonal variation in household electricity demand: A '
 'comparison of monitored and synthetic daily load profiles. <i>Energy and Buildings, 179</i>, '
 '292–300. https://doi.org/10.1016/j.enbuild.2018.09.030',
 'National Oceanic and Atmospheric Administration. (2026). <i>nClimDiv statewide cooling degree '
 'days, heating degree days, and average temperature</i> (Version 1.0.0, 2026-06-04) [Data set]. '
 'National Centers for Environmental Information. https://www.ncei.noaa.gov/pub/data/cirs/climdiv/',
 'Newey, W. K., & West, K. D. (1987). A simple, positive semi-definite, heteroskedasticity and '
 'autocorrelation consistent covariance matrix. <i>Econometrica, 55</i>(3), 703–708. '
 'https://doi.org/10.2307/1913610',
 'U.S. Energy Information Administration. (2022, May 31). Nearly 90% of U.S. households used air '
 'conditioning in 2020. <i>Today in Energy.</i> https://www.eia.gov/todayinenergy/detail.php?id=52558',
 'U.S. Energy Information Administration. (2023a). <i>Highlights for space heating in U.S. homes by '
 'state, 2020</i> [Data tables]. Residential Energy Consumption Survey. '
 'https://www.eia.gov/consumption/residential/data/2020/state/pdf/State%20Space%20Heating.pdf',
 'U.S. Energy Information Administration. (2023b). <i>Units and calculators explained: Degree '
 'days.</i> https://www.eia.gov/energyexplained/units-and-calculators/degree-days.php',
 'U.S. Energy Information Administration. (2024, March 15). <i>How much electricity is used for air '
 'conditioning in the United States?</i> Frequently Asked Questions. '
 'https://www.eia.gov/tools/faqs/faq.php?id=1174&t=1',
 'U.S. Energy Information Administration. (2026). <i>Form EIA-861M: Monthly electric power industry '
 'report</i> [Data set]. https://www.eia.gov/electricity/data/eia861m/',
]
for ref in refs:
    q = doc.add_paragraph()
    q.paragraph_format.line_spacing = 2.0
    q.paragraph_format.left_indent = Inches(0.5)
    q.paragraph_format.first_line_indent = Inches(-0.5)
    q.paragraph_format.space_after = Pt(0)
    for chunk in re.split(r'(<i>.*?</i>)', ref):
        if not chunk: continue
        if chunk.startswith('<i>'): q.add_run(chunk[3:-4]).italic = True
        else: q.add_run(chunk)

pagebreak()
h1('Appendix')
para('Residual Diagnostics', bold=True, indent=False)
para(
    'This appendix presents residual diagnostics for the cooling-only baseline (Model 1) and for '
    'the two specifications that carry the substantive results (Models 3 and 4), together with a '
    'normality check for the paired comparison.')
for n, key, path, extra in [
        ('A1', 'Model 1 (Cooling Degree Days Only)', f'{FIG}A1_diag_m1.png',
         f'The lag-1 residual autocorrelation is {nz(R["acf1_M1"])} and the seasonal value at lag 12 '
         f'is {nz(R["acf12_M1"])}. The dense column in Panel A contains the zero-CDD months, which the '
         'model maps to a single fitted value.'),
        ('A2', 'Model 3 (Cooling and Heating Degree Days With a Linear Trend)', f'{FIG}A2_diag_m3.png',
         f'The lag-1 value falls to {nz(R["acf1_M3"])} and Durbin–Watson rises to '
         f'{num(M["M3"]["dw"])}, but a seasonal component of {nz(R["acf12_M3"])} remains at lag 12.'),
        ('A3', 'Model 4 (Model 3 Plus Calendar-Month Fixed Effects)', f'{FIG}A3_diag_m4.png',
         f'Calendar-month effects reduce the lag-12 value to {nz(R["acf12_M4"])}. Residual '
         'dependence is reduced but not eliminated in any specification.')]:
    figure(n, f'Residual Diagnostics for {key}', path,
           'Panel A plots residuals against fitted values; Panel B the residual distribution with a '
           'matched normal curve; Panel C the normal quantile–quantile plot; Panel D the residual '
           'autocorrelation function with 95% bounds for white noise. ' + extra, width=6.2)
figure('A4', 'Normal Q–Q Plot of the July–January Paired Differences', f'{FIG}A4_paired_qq.png',
       'The line marks agreement with a normal distribution, fitted through the first and third '
       'quartiles. No marked departure is evident, supporting the normality assumption of the '
       'paired t test and the equivalence tests. This plot assesses the distribution of the '
       'differences, not their independence, which is addressed in the text.', width=4.3)

doc.save(OUT)
print('saved:', OUT, os.path.getsize(OUT), 'bytes')
